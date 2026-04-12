#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
用python-pptx画小房子
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RgbColor
from pptx.enum.shapes import MSO_SHAPE

def draw_house():
    # 创建演示文稿
    prs = Presentation()
    
    # 设置幻灯片大小为标准4:3
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    # 添加空白幻灯片
    slide_layout = prs.slide_layouts[6]  # 空白布局
    slide = prs.slides.add_slide(slide_layout)
    
    # ===== 画地面（草坪）=====
    ground = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(5.5),
        Inches(10), Inches(2)
    )
    ground.fill.solid()
    ground.fill.fore_color.rgb = RgbColor(34, 139, 34)  # 绿色
    ground.line.fill.background()
    
    # ===== 画房子主体（墙壁）=====
    house_body = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(3), Inches(3),
        Inches(4), Inches(2.5)
    )
    house_body.fill.solid()
    house_body.fill.fore_color.rgb = RgbColor(255, 228, 196)  # 米色
    house_body.line.fill.background()
    
    # ===== 画屋顶（三角形）=====
    roof = slide.shapes.add_shape(
        MSO_SHAPE.ISOSCELES_TRIANGLE,
        Inches(2.5), Inches(1.5),
        Inches(5), Inches(1.8)
    )
    roof.fill.solid()
    roof.fill.fore_color.rgb = RgbColor(139, 69, 19)  # 棕色
    roof.line.fill.background()
    
    # ===== 画烟囱 ======
    chimney = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(5.5), Inches(1.3),
        Inches(0.6), Inches(1)
    )
    chimney.fill.solid()
    chimney.fill.fore_color.rgb = RgbColor(139, 69, 19)  # 棕色
    chimney.line.fill.background()
    
    # ===== 画烟囱冒出的烟 ======
    # 烟圈1
    smoke1 = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(5.65), Inches(0.5),
        Inches(0.4), Inches(0.3)
    )
    smoke1.fill.solid()
    smoke1.fill.fore_color.rgb = RgbColor(200, 200, 200)  # 灰色
    smoke1.line.fill.background()
    
    # 烟圈2
    smoke2 = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(5.75), Inches(0.1),
        Inches(0.5), Inches(0.4)
    )
    smoke2.fill.solid()
    smoke2.fill.fore_color.rgb = RgbColor(180, 180, 180)  # 灰色
    smoke2.line.fill.background()
    
    # ===== 画门 ======
    door = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(4.3), Inches(4),
        Inches(0.8), Inches(1.5)
    )
    door.fill.solid()
    door.fill.fore_color.rgb = RgbColor(101, 67, 33)  # 深棕色
    door.line.fill.background()
    
    # 门把手
    doorknob = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(4.85), Inches(4.8),
        Inches(0.12), Inches(0.12)
    )
    doorknob.fill.solid()
    doorknob.fill.fore_color.rgb = RgbColor(255, 215, 0)  # 金色
    doorknob.line.fill.background()
    
    # ===== 画窗户 ======
    # 左窗户
    left_window = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(3.3), Inches(3.5),
        Inches(0.8), Inches(0.8)
    )
    left_window.fill.solid()
    left_window.fill.fore_color.rgb = RgbColor(135, 206, 250)  # 浅蓝色
    left_window.line.fill.background()
    
    # 右窗户
    right_window = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(5.9), Inches(3.5),
        Inches(0.8), Inches(0.8)
    )
    right_window.fill.solid()
    right_window.fill.fore_color.rgb = RgbColor(135, 206, 250)  # 浅蓝色
    right_window.line.fill.background()
    
    # ===== 画太阳 ======
    sun = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(8), Inches(1),
        Inches(1.2), Inches(1.2)
    )
    sun.fill.solid()
    sun.fill.fore_color.rgb = RgbColor(255, 255, 0)  # 黄色
    sun.line.fill.background()
    
    # ===== 画云朵 ======
    # 云朵1
    cloud1 = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(1), Inches(1.2),
        Inches(1.2), Inches(0.6)
    )
    cloud1.fill.solid()
    cloud1.fill.fore_color.rgb = RgbColor(255, 255, 255)  # 白色
    cloud1.line.fill.background()
    
    # 云朵2
    cloud2 = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(1.5), Inches(1),
        Inches(1), Inches(0.5)
    )
    cloud2.fill.solid()
    cloud2.fill.fore_color.rgb = RgbColor(255, 255, 255)  # 白色
    cloud2.line.fill.background()
    
    # ===== 画小路 ======
    path = slide.shapes.add_shape(
        MSO_SHAPE.TRAPEZOID,
        Inches(4.3), Inches(5.5),
        Inches(0.8), Inches(1.5)
    )
    path.fill.solid()
    path.fill.fore_color.rgb = RgbColor(210, 180, 140)  # 土黄色
    path.line.fill.background()
    
    # ===== 画树 ======
    # 树干
    tree_trunk = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(1), Inches(4.2),
        Inches(0.5), Inches(1.3)
    )
    tree_trunk.fill.solid()
    tree_trunk.fill.fore_color.rgb = RgbColor(139, 69, 19)  # 棕色
    tree_trunk.line.fill.background()
    
    # 树冠
    tree_leaves = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(0.5), Inches(2.8),
        Inches(1.5), Inches(1.6)
    )
    tree_leaves.fill.solid()
    tree_leaves.fill.fore_color.rgb = RgbColor(0, 128, 0)  # 深绿色
    tree_leaves.line.fill.background()
    
    # 保存文件
    prs.save('小房子.pptx')
    print("✅ 小房子PPT已创建成功！文件名为：小房子.pptx")

if __name__ == "__main__":
    draw_house()
